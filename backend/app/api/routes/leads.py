from fastapi import APIRouter, Depends, HTTPException, Response, Request
from fastapi.responses import FileResponse
from typing import List, Optional
from datetime import datetime, timezone
import os
import time
import json
from bson import ObjectId
import io
import openpyxl
from pydantic import BaseModel
from motor.motor_asyncio import AsyncIOMotorDatabase
from app.db.database import get_db
from app.api.dependencies import get_current_user, RequireRole
from app.models.user import UserInDB
from app.models.lead import LeadCreate, LeadResponse, PaginatedLeadResponse, build_lead_response
from app.core.logger import logger

router = APIRouter()

async def _reconcile_lead_ai_status(lead: dict, db: AsyncIOMotorDatabase) -> dict:
    """Self-healing helper: ensure lead ai_analysis state accurately reflects genuine worker jobs.
    Prevents leads from being orphaned in infinite 'Queued' or 'Analyzing...' state,
    and ensures active/completed/failed jobs in ai_analysis_jobs are reflected on the lead."""
    if not isinstance(lead, dict):
        return lead

    ai = lead.get("ai_analysis")
    if not isinstance(ai, dict):
        ai = {}
        lead["ai_analysis"] = ai

    raw_status = str(ai.get("analysis_status") or ai.get("status") or "").lower()
    lead_id_str = str(lead.get("_id", ""))
    if not lead_id_str:
        return lead

    id_filters = [lead_id_str]
    if ObjectId.is_valid(lead_id_str):
        id_filters.append(ObjectId(lead_id_str))

    has_score = bool(
        (lead.get("ai_score") and lead.get("ai_score") > 0) or
        (ai.get("overall_score") and ai.get("overall_score") > 0) or
        (ai.get("overall_ai_score") and ai.get("overall_ai_score") > 0) or
        (lead.get("lead_score") and lead.get("lead_score") > 0)
    )

    now = datetime.now(timezone.utc)

    # 1. Query the latest job associated with this lead (direct or linked in bulk)
    job_query = {
        "$or": [
            {"lead_id": {"$in": id_filters}},
            {"linked_lead_ids": lead_id_str}
        ]
    }

    try:
        latest_job = await db["ai_analysis_jobs"].find_one(job_query, sort=[("_id", -1)])
    except Exception as job_err:
        logger.warning(f"Failed to lookup ai_analysis_jobs for lead {lead_id_str}: {job_err}")
        latest_job = None

    if latest_job:
        job_status = str(latest_job.get("status") or "").upper()
        active_statuses = {"WAITING", "QUEUED", "PENDING", "RUNNING", "STARTED", "COLLECTING", "ANALYZING", "SCORING", "SAVING", "RETRY"}

        if job_status in active_statuses:
            job_updated = latest_job.get("updated_at") or latest_job.get("started_at") or latest_job.get("created_at")
            if job_updated:
                if isinstance(job_updated, str):
                    try:
                        job_updated = datetime.fromisoformat(job_updated.replace('Z', '+00:00'))
                    except Exception:
                        job_updated = None
                if isinstance(job_updated, datetime) and job_updated.tzinfo is None:
                    job_updated = job_updated.replace(tzinfo=timezone.utc)
                age_seconds = (now - job_updated).total_seconds() if isinstance(job_updated, datetime) else 0
            else:
                age_seconds = 0

            if age_seconds > 300:
                # Active job is stale/hung beyond 5 minutes -> fail it
                await db["ai_analysis_jobs"].update_one(
                    {"_id": latest_job["_id"]},
                    {"$set": {"status": "FAILED", "error_message": "Analysis timed out after 300s.", "updated_at": now}}
                )
                new_status = "Completed" if has_score else "Failed"
                stage_label = "Completed" if has_score else "Analysis timed out. Please retry."
                prog_val = 100 if has_score else 0
                ai["analysis_status"] = new_status
                ai["status"] = new_status
                ai["stage"] = stage_label
                ai["progress"] = prog_val
                await db["leads"].update_one(
                    {"_id": lead["_id"]},
                    {"$set": {
                        "ai_analysis.analysis_status": new_status,
                        "ai_analysis.status": new_status,
                        "ai_analysis.stage": stage_label,
                        "ai_analysis.progress": prog_val,
                        "updated_at": now,
                    }}
                )
            else:
                # Sync real-time progress and stage from active job
                canonical_status = "Analyzing" if job_status in ["STARTED", "COLLECTING", "ANALYZING", "SCORING", "SAVING", "RUNNING"] else "Queued"
                ai["analysis_status"] = canonical_status
                ai["status"] = canonical_status
                ai["progress"] = latest_job.get("progress") if latest_job.get("progress") is not None else (25 if canonical_status == "Analyzing" else 10)
                ai["stage"] = latest_job.get("current_stage") or ("Analyzing website..." if canonical_status == "Analyzing" else "Queued for AI analysis")
                ai["queue_id"] = str(latest_job.get("queue_id") or "")
                started_dt = latest_job.get("started_at") or latest_job.get("created_at") or now
                ai["analysis_started_at"] = started_dt

        elif job_status in {"COMPLETED", "SUCCESS", "PARTIAL_SUCCESS"}:
            if raw_status not in {"completed", "done", "partial_success", "success"}:
                ai["analysis_status"] = "Completed"
                ai["status"] = "Completed"
                ai["progress"] = 100
                ai["stage"] = "Completed"
                if latest_job.get("overall_ai_score") and not lead.get("ai_score"):
                    lead["ai_score"] = latest_job["overall_ai_score"]
                    ai["overall_ai_score"] = latest_job["overall_ai_score"]

        elif job_status in {"FAILED", "CANCELLED", "TIMEOUT"}:
            if raw_status not in {"completed", "done", "partial_success", "success", "failed"}:
                ai["analysis_status"] = "Failed"
                ai["status"] = "Failed"
                ai["progress"] = 0
                ai["stage"] = latest_job.get("error_message") or "Analysis failed. Please retry."
                await db["leads"].update_one(
                    {"_id": lead["_id"]},
                    {"$set": {
                        "ai_analysis.analysis_status": "Failed",
                        "ai_analysis.status": "Failed",
                        "ai_analysis.stage": ai["stage"],
                        "ai_analysis.progress": 0,
                        "updated_at": now,
                    }}
                )

    else:
        # No job record found in ai_analysis_jobs
        if has_score:
            ai["analysis_status"] = "Completed"
            ai["status"] = "Completed"
            ai["progress"] = 100
            ai["stage"] = "Completed"
        elif raw_status in {"queued", "running", "analyzing", "processing", "in_progress", "started", "collecting", "scoring", "saving", "waiting", "pending"}:
            started_at = ai.get("analysis_started_at") or ai.get("started_at")
            age_seconds = 0
            if started_at:
                if isinstance(started_at, str):
                    try:
                        started_at = datetime.fromisoformat(started_at.replace('Z', '+00:00'))
                    except Exception:
                        started_at = None
                if isinstance(started_at, datetime):
                    if started_at.tzinfo is None:
                        started_at = started_at.replace(tzinfo=timezone.utc)
                    age_seconds = (now - started_at).total_seconds()

            queue_id = ai.get("queue_id")
            has_active_queue = False
            if queue_id:
                q_obj = ObjectId(queue_id) if ObjectId.is_valid(str(queue_id)) else str(queue_id)
                q = await db["ai_analysis_queues"].find_one({"_id": q_obj, "status": {"$in": ["WAITING", "RUNNING", "STARTED", "PENDING"]}})
                if q:
                    q_updated = q.get("updated_at") or q.get("created_at")
                    if isinstance(q_updated, str):
                        try:
                            q_updated = datetime.fromisoformat(q_updated.replace('Z', '+00:00'))
                        except Exception:
                            q_updated = None
                    if isinstance(q_updated, datetime) and q_updated.tzinfo is None:
                        q_updated = q_updated.replace(tzinfo=timezone.utc)
                    q_age = (now - q_updated).total_seconds() if isinstance(q_updated, datetime) else 999
                    has_active_queue = q_age <= 300

            if (not started_at or age_seconds > 300) and not has_active_queue:
                ai["analysis_status"] = "Failed"
                ai["status"] = "Failed"
                ai["stage"] = "Analysis incomplete. Please retry."
                ai["progress"] = 0
                await db["leads"].update_one(
                    {"_id": lead["_id"]},
                    {"$set": {
                        "ai_analysis.analysis_status": "Failed",
                        "ai_analysis.status": "Failed",
                        "ai_analysis.stage": "Analysis incomplete. Please retry.",
                        "ai_analysis.progress": 0,
                        "updated_at": now,
                    }}
                )
            else:
                ai["analysis_status"] = "Queued"
                ai["status"] = "Queued"
                if not ai.get("progress"):
                    ai["progress"] = 10
                if not ai.get("stage"):
                    ai["stage"] = "Queued for AI analysis"
        else:
            if not ai.get("analysis_status"):
                ai["analysis_status"] = "not_analyzed"
                ai["status"] = "not_analyzed"
                ai["progress"] = 0
                ai["stage"] = "Not Analyzed"

    return lead

class DBInfoResponse(BaseModel):
    total_leads: int
    db_name: str
    collection_name: str
    estimated_size: str
    current_time: str

class BulkIds(BaseModel):
    lead_ids: List[str]

class LeadIdsResponse(BaseModel):
    data: List[str]
    total: int


def build_lead_filters(
    search_term: Optional[str] = None,
    job_id: Optional[str] = None,
    city: Optional[str] = None,
    category: Optional[str] = None,
    has_website: Optional[bool] = None,
    has_phone: Optional[bool] = None,
    has_email: Optional[bool] = None,
    min_rating: Optional[float] = None,
    status: Optional[str] = None,
) -> dict:
    """Build the Mongo query used to filter leads on the /leads page."""
    query: dict = {}

    # Combine independent filters with $and so each one keeps its own shape.
    filter_clauses = []
    if search_term:
        filter_clauses.append({
            "$or": [
                {"name": {"$regex": search_term, "$options": "i"}},
                {"address": {"$regex": search_term, "$options": "i"}},
                {"phone": {"$regex": search_term, "$options": "i"}},
                {"email": {"$regex": search_term, "$options": "i"}}
            ]
        })
    if job_id:
        # Join Leads to a Search Session using every existing relationship
        # field (source_job_id / job_id / search_job_id). Never filters by
        # keyword because keywords can repeat across sessions.
        filter_clauses.append({
            "$or": [
                {"source_job_id": job_id},
                {"job_id": job_id},
                {"search_job_id": job_id},
            ]
        })
    if filter_clauses:
        query["$and"] = filter_clauses

    # City filter
    if city:
        query["city"] = {"$regex": city, "$options": "i"}

    # Category filter
    if category:
        query["category"] = {"$regex": category, "$options": "i"}

    if has_website is True:
        query["website"] = {"$nin": [None, ""]}
    elif has_website is False:
        query["website"] = {"$in": [None, ""]}

    if has_phone is True:
        query["phone"] = {"$nin": [None, ""]}
    elif has_phone is False:
        query["phone"] = {"$in": [None, ""]}

    if has_email is True:
        query["email"] = {"$nin": [None, ""]}
    elif has_email is False:
        query["email"] = {"$in": [None, ""]}

    if min_rating is not None:
        query["google_rating"] = {"$gte": min_rating}

    if status and status != "ALL":
        query["status"] = status

    return query

class FavoriteUpdate(BaseModel):
    is_favorite: bool

class TagsUpdate(BaseModel):
    tags: List[str]
    
class NoteCreate(BaseModel):
    lead_id: str
    content: str

# Register both "" and "/" so GET /api/v1/leads (no trailing slash) resolves
# directly instead of issuing a 307 redirect. Behind nginx the redirect gets
# rewritten to /api/v1/api/v1/leads/ (double prefix) and 404s, which made the
# Leads page always render its empty state. This mirrors the search route.
@router.get("", response_model=PaginatedLeadResponse)
@router.get("/", response_model=PaginatedLeadResponse, include_in_schema=False)
async def get_leads(
    limit: int = 10,
    skip: int = 0,
    search: Optional[str] = None,
    keyword: Optional[str] = None,  # Alias for search
    job_id: Optional[str] = None,  # Filter leads belonging to a specific search session
    city: Optional[str] = None,
    category: Optional[str] = None,
    has_website: Optional[bool] = None,
    has_phone: Optional[bool] = None,
    has_email: Optional[bool] = None,
    min_rating: Optional[float] = None,
    status: Optional[str] = None,
    sort_by: str = "created_at",
    sort_order: int = -1,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    # Use keyword if provided, otherwise use search
    search_term = keyword or search

    query = build_lead_filters(
        search_term=search_term,
        job_id=job_id,
        city=city,
        category=category,
        has_website=has_website,
        has_phone=has_phone,
        has_email=has_email,
        min_rating=min_rating,
        status=status,
    )
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id

    # Validation of sort fields
    if sort_by not in ["name", "google_rating", "review_count", "created_at", "category", "status", "lead_score", "ai_score"]:
        sort_by = "created_at"

    raw_leads = await cursor.to_list(length=limit)
    reconciled_leads = await asyncio.gather(*[_reconcile_lead_ai_status(l, db) for l in raw_leads])

    leads = []
    warnings: list = []
    phones_returned = 0
    locations_returned = 0
    for lead in reconciled_leads:
        try:
            normalized, lead_warnings = build_lead_response(lead, logger=logger)
            if lead_warnings:
                warnings.extend(lead_warnings)
                for w in lead_warnings:
                    logger.warning(f"LEAD_NORMALIZATION: {w}")
            if getattr(normalized, "primary_phone", None) or getattr(normalized, "phone", None):
                phones_returned += 1
            if getattr(normalized, "state", None) or getattr(normalized, "city", None) or getattr(normalized, "area", None):
                locations_returned += 1
            leads.append(normalized)
        except Exception as e:
            logger.error(f"LEAD_VALIDATION_SKIP: Skipping lead {lead.get('_id')} ({lead.get('name')}): {e}")
            continue

    total = await db["leads"].count_documents(query)
    logger.info(
        f"PHONES_RETURNED_BY_API: page={skip//max(limit,1)} returned={len(leads)} "
        f"leads_with_phone={phones_returned} leads_with_location={locations_returned}"
    )
    return PaginatedLeadResponse(data=leads, total=total, warnings=warnings)

@router.get("/export")
async def export_leads(
    format: str = "excel",
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    query = {}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    cursor = db["leads"].find(query)
    leads = await cursor.to_list(length=None)
    
    # Fetch job keywords to populate Category column with module name
    job_ids = list(set(
        lead.get("job_id") or lead.get("search_job_id") or lead.get("source_job_id")
        for lead in leads if (lead.get("job_id") or lead.get("search_job_id") or lead.get("source_job_id"))
    ))
    jobs = await db["search_jobs"].find({"_id": {"$in": [ObjectId(jid) for jid in job_ids if ObjectId.is_valid(jid)]}}).to_list(length=None)
    job_keywords = {str(job["_id"]): job.get("keyword", "") for job in jobs}
    
    headers = [
        "Lead Name",
        "Lead Category (Searched Keyword)",
        "Phone Number",
        "Website",
        "Country",
        "State",
        "Address",
    ]

    def _get_phone(lead: dict) -> str:
        for key in ("phone", "primary_phone", "whatsapp_number"):
            value = lead.get(key)
            if value:
                return str(value).strip()
        for key in ("phones", "all_phones", "contact_discovery"):
            value = lead.get(key)
            if isinstance(value, list):
                for p in value:
                    if p:
                        return str(p).strip()
            elif isinstance(value, dict):
                for p in (value.get("primary_phone"), value.get("secondary_phones") or []):
                    if isinstance(p, list):
                        for item in p:
                            if item:
                                return str(item).strip()
                    elif p:
                        return str(p).strip()
        return ""

    def _build_row(lead: dict) -> list:
        job_id = str(lead.get("job_id") or lead.get("search_job_id") or lead.get("source_job_id") or "")
        module_name = (
            job_keywords.get(job_id)
            or lead.get("search_keyword")
            or lead.get("category")
            or ""
        )
        country = lead.get("country") or (lead.get("website_information") or {}).get("country") or ""
        return [
            lead.get("name") or lead.get("business_name") or "",
            module_name,
            _get_phone(lead),
            lead.get("website") or "",
            country,
            lead.get("state") or "",
            lead.get("address") or "",
        ]
    
    if format.lower() == "csv":
        import csv
        stream = io.StringIO()
        writer = csv.writer(stream)
        writer.writerow(headers)
        
        for lead in leads:
            writer.writerow(_build_row(lead))
            
        content = stream.getvalue()
        response_headers = {
            "Content-Disposition": "attachment; filename=leads.csv"
        }
        return Response(content=content, media_type="text/csv", headers=response_headers)
        
    else:  # excel / xlsx / xls
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        
        for lead in leads:
            ws.append(_build_row(lead))
            
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        
        response_headers = {
            "Content-Disposition": "attachment; filename=leads.xlsx"
        }
        return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=response_headers)

@router.post("", response_model=LeadResponse)
@router.post("/", response_model=LeadResponse, include_in_schema=False)
async def create_lead(
    lead_in: LeadCreate,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    lead_dict = lead_in.model_dump(by_alias=True)
    lead_dict["user_id"] = current_user.id
    lead_dict["created_at"] = datetime.utcnow()
    lead_dict["updated_at"] = datetime.utcnow()
    if "ai_analysis" not in lead_dict:
        lead_dict["ai_analysis"] = {
            "analysis_status": "not_analyzed",
            "progress": 0,
            "stage": "Not Analyzed"
        }
    result = await db["leads"].insert_one(lead_dict)
    lead_dict["_id"] = str(result.inserted_id)
    resp, _ = build_lead_response(lead_dict, logger=logger)
    return resp

@router.post("/bulk-delete")
async def bulk_delete(
    data: BulkIds,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    valid_ids = [ObjectId(lid) for lid in data.lead_ids if ObjectId.is_valid(lid)]
    await db["leads"].delete_many({"_id": {"$in": valid_ids}, "user_id": current_user.id})
    return {"success": True}

@router.get("/ids", response_model=LeadIdsResponse)
async def get_lead_ids(
    keyword: Optional[str] = None,
    job_id: Optional[str] = None,
    city: Optional[str] = None,
    category: Optional[str] = None,
    has_website: Optional[bool] = None,
    has_phone: Optional[bool] = None,
    has_email: Optional[bool] = None,
    min_rating: Optional[float] = None,
    status: Optional[str] = None,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    """Return every lead _id matching the same filters as GET /leads so the
    Leads page can "Select All" across all pages at once."""
    query = build_lead_filters(
        search_term=keyword,
        job_id=job_id,
        city=city,
        category=category,
        has_website=has_website,
        has_phone=has_phone,
        has_email=has_email,
        min_rating=min_rating,
        status=status,
    )
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    cursor = db["leads"].find(query, {"_id": 1})
    ids = []
    async for doc in cursor:
        ids.append(str(doc["_id"]))
    return LeadIdsResponse(data=ids, total=len(ids))

@router.get("/db-info", response_model=DBInfoResponse)
async def get_db_info(
    current_user: UserInDB = Depends(RequireRole(["ADMIN"])),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    total_leads = await db["leads"].count_documents({})
    # Size estimate: openpyxl spreadsheet size is roughly 1.2KB per lead.
    size_kb = round((total_leads * 1.2), 2)
    size_str = f"{size_kb} KB" if size_kb < 1024 else f"{round(size_kb / 1024, 2)} MB"
    
    return DBInfoResponse(
        total_leads=total_leads,
        db_name=db.name,
        collection_name="leads",
        estimated_size=size_str,
        current_time=datetime.utcnow().isoformat() + "Z"
    )

@router.post("/export-and-delete")
async def export_and_delete_leads(
    request: Request,
    current_user: UserInDB = Depends(RequireRole(["ADMIN"])),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    logger.info("Admin Delete All Leads: Export and delete process started.")
    start_time = time.time()
    
    # 1. Fetch all leads
    cursor = db["leads"].find({})
    leads = await cursor.to_list(length=None)
    total_count = len(leads)
    
    # Fetch job keywords to populate Category column with module name
    job_ids = list(set(
        lead.get("job_id") or lead.get("search_job_id") or lead.get("source_job_id")
        for lead in leads if (lead.get("job_id") or lead.get("search_job_id") or lead.get("source_job_id"))
    ))
    jobs = await db["search_jobs"].find({"_id": {"$in": [ObjectId(jid) for jid in job_ids if ObjectId.is_valid(jid)]}}).to_list(length=None)
    job_keywords = {str(job["_id"]): job.get("keyword", "") for job in jobs}
    
    # 2. Create backup folder if missing
    backup_dir = os.path.abspath(os.path.join(os.getcwd(), "backups", "leads"))
    os.makedirs(backup_dir, exist_ok=True)
    
    # Generate timestamped backup filename
    timestamp = datetime.utcnow().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"Lead_Backup_{timestamp}.xlsx"
    backup_filepath = os.path.join(backup_dir, filename)
    
    logger.info(f"Backup Started: Generating backup Excel file '{filename}'...")
    
    # 3. Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads Backup"
    
    # Write headers
    headers = [
        "Business Name", "Category", "Phone", "Email", "Website", "Address", 
        "City", "State", "Country", "Latitude", "Longitude", "Google Rating", 
        "Reviews", "Lead Score", "AI Score", "Tags", "Created Date", 
        "Updated Date", "Google Maps URL", "Notes", "Owner", "Social Links", 
        "Technology", "Place ID", "Opening Hours", "Business Status"
    ]
    ws.append(headers)
    
    # Append lead rows
    for lead in leads:
        lead_id_str = str(lead.get("_id", ""))
        job_id = str(lead.get("job_id") or lead.get("search_job_id") or lead.get("source_job_id") or "")
        module_name = job_keywords.get(job_id) or lead.get("search_keyword") or lead.get("category", "")
        
        # Get lead notes
        notes_cursor = db["notes"].find({"lead_id": lead_id_str})
        notes_list = await notes_cursor.to_list(length=None)
        notes_str = " | ".join([n.get("content", "").replace("\n", " ") for n in notes_list if n.get("content")])
        
        # Format tags
        tags = lead.get("tags", [])
        tags_str = ", ".join(tags) if isinstance(tags, list) else str(tags)
        
        # Format created/updated dates
        created_at = lead.get("created_at")
        created_str = created_at.isoformat() if isinstance(created_at, datetime) else str(created_at or "")
        updated_at = lead.get("updated_at")
        updated_str = updated_at.isoformat() if isinstance(updated_at, datetime) else str(updated_at or "")
        
        # Format social links & technologies
        social_links = lead.get("social_links", {})
        social_str = json.dumps(social_links) if isinstance(social_links, dict) else str(social_links or "")
        technologies = lead.get("technologies", [])
        tech_str = ", ".join(technologies) if isinstance(technologies, list) else str(technologies or "")
        
        # Format opening_hours
        opening_hours = lead.get("opening_hours", [])
        if isinstance(opening_hours, list):
            hours_str = " | ".join(str(h) for h in opening_hours)
        else:
            hours_str = str(opening_hours or "")
            
        # Ensure email is string
        email = lead.get("email", "")
        email_str = ", ".join(email) if isinstance(email, list) else str(email or "")

        row = [
            lead.get("name", ""),
            module_name,
            lead.get("phone", ""),
            email_str,
            lead.get("website", ""),
            lead.get("address", ""),
            lead.get("city", ""),
            lead.get("state", ""),
            lead.get("country", ""),
            lead.get("latitude", ""),
            lead.get("longitude", ""),
            lead.get("google_rating", ""),
            lead.get("review_count", 0),
            lead.get("lead_score", 0),
            lead.get("website_score", 0),
            tags_str,
            created_str,
            updated_str,
            lead.get("source_url", ""),
            notes_str,
            lead.get("user_id", ""),
            social_str,
            tech_str,
            lead.get("place_id", ""),
            hours_str,
            lead.get("business_status", "")
        ]
        ws.append(row)
        
    wb.save(backup_filepath)
    logger.info(f"Backup Completed: Excel file saved to {backup_filepath}")
    
    # 4. Verify backup file
    if not os.path.exists(backup_filepath) or os.path.getsize(backup_filepath) == 0:
        logger.error(f"Backup Verification Failed: File {backup_filepath} not found or empty.")
        raise HTTPException(status_code=500, detail="Backup generation failed: file empty or missing.")
        
    wb_verify = openpyxl.load_workbook(backup_filepath, read_only=True)
    ws_verify = wb_verify.active
    row_count = ws_verify.max_row
    wb_verify.close()
    
    # Total rows should be total_count + 1 (headers)
    if row_count - 1 != total_count:
        logger.error(f"Backup Verification Failed: Row count mismatch. Document count: {total_count}, Excel rows (minus header): {row_count - 1}")
        raise HTTPException(
            status_code=500, 
            detail=f"Backup verification failed: document count mismatch. DB: {total_count}, Backup: {row_count - 1}"
        )
        
    logger.info("Backup Verified: Excel row count matches MongoDB document count exactly.")
    
    # 5. Perform deletion
    logger.info("Deletion Started: Clearing leads collection and search history...")
    delete_result = await db["leads"].delete_many({})
    deleted_count = delete_result.deleted_count
    
    # Clear search history collections
    await db["search_jobs"].delete_many({})
    await db["search_queue"].delete_many({})
    
    logger.info(f"Deletion Completed: Successfully deleted {deleted_count} documents from leads collection and cleared search history.")
    
    # 6. Save audit log
    logger.info("Audit Saved: Writing deletion details to audit_logs collection.")
    audit_log = {
        "admin_id": current_user.id,
        "admin_email": current_user.email,
        "ip_address": request.client.host,
        "action": "DELETE_ALL_LEADS",
        "deleted_count": deleted_count,
        "backup_filename": filename,
        "backup_path": backup_filepath,
        "timestamp": datetime.utcnow()
    }
    await db["audit_logs"].insert_one(audit_log)
    
    execution_time_ms = int((time.time() - start_time) * 1000)
    
    return {
        "success": True,
        "deleted_count": deleted_count,
        "backup_file": filename,
        "backup_path": backup_filepath,
        "execution_time": f"{execution_time_ms}ms"
    }

@router.get("/download-backup")
async def download_backup(
    filename: str,
    current_user: UserInDB = Depends(RequireRole(["ADMIN"])),
):
    backup_dir = os.path.abspath(os.path.join(os.getcwd(), "backups", "leads"))
    target_path = os.path.abspath(os.path.join(backup_dir, filename))
    
    # Path traversal vulnerability prevention
    if not target_path.startswith(backup_dir):
        raise HTTPException(status_code=400, detail="Invalid filename parameter")
        
    if not os.path.exists(target_path):
        raise HTTPException(status_code=404, detail="Backup file not found")
        
    return FileResponse(
        target_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename
    )

@router.get("/{id}", response_model=LeadResponse)
async def get_lead(
    id: str,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    if not ObjectId.is_valid(id): raise HTTPException(400)
    query = {"_id": ObjectId(id)}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    lead = await db["leads"].find_one(query)
    if not lead:
        raise HTTPException(404, detail="Lead not found")

    lead = await _reconcile_lead_ai_status(lead, db)
    lead["_id"] = str(lead["_id"])
    try:
        normalized, warnings = build_lead_response(lead, logger=logger)
        for w in warnings:
            logger.warning(f"LEAD_NORMALIZATION: {w}")
        return normalized
    except Exception as e:
        logger.error(f"LEAD_RECOVERY_FAILED: Could not normalize lead {lead.get('_id')}: {e}")
        # Structured fallback: never return 500 for malformed data.
        return LeadResponse(**{k: v for k, v in lead.items() if k in LeadResponse.model_fields})

@router.delete("/{id}")
async def delete_lead(
    id: str,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    if not ObjectId.is_valid(id): raise HTTPException(400)
    query = {"_id": ObjectId(id)}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    await db["leads"].delete_one(query)
    return {"success": True}

@router.patch("/{id}/favorite")
async def patch_favorite(
    id: str,
    data: FavoriteUpdate,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    if not ObjectId.is_valid(id): raise HTTPException(400)
    query = {"_id": ObjectId(id)}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    await db["leads"].update_one(
        query,
        {"$set": {"is_favorite": data.is_favorite}}
    )
    return {"success": True}

@router.patch("/{id}/tags")
async def patch_tags(
    id: str,
    data: TagsUpdate,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    if not ObjectId.is_valid(id): raise HTTPException(400)
    query = {"_id": ObjectId(id)}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    await db["leads"].update_one(
        query,
        {"$set": {"tags": data.tags}}
    )
    return {"success": True}

@router.get("/{id}/notes")
async def get_lead_notes(
    id: str,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    query = {"lead_id": id}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    cursor = db["notes"].find(query).sort("created_at", -1)
    notes = await cursor.to_list(length=None)
    for n in notes: n["_id"] = str(n["_id"])
    return notes

@router.post("/{id}/notes")
async def create_lead_note(
    id: str,
    data: NoteCreate,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    note = {
        "lead_id": id,
        "user_id": current_user.id,
        "content": data.content,
        "created_at": datetime.utcnow()
    }
    await db["notes"].insert_one(note)
    return {"success": True}

class StartAnalysisResponse(BaseModel):
    lead_id: str
    website_url: str
    queue_id: str
    worker_id: str

@router.post("/{id}/start-analysis", response_model=StartAnalysisResponse)
async def start_analysis(
    id: str,
    current_user: UserInDB = Depends(get_current_user),
    db: AsyncIOMotorDatabase = Depends(get_db)
):
    if not ObjectId.is_valid(id):
        raise HTTPException(status_code=400, detail="Invalid lead ID format.")
    
    query = {"_id": ObjectId(id)}
    if current_user.role != "ADMIN":
        query["user_id"] = current_user.id
    lead = await db["leads"].find_one(query)
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found.")

    website = str(lead.get("website", "")).strip()

    import re
    from urllib.parse import urlparse

    def _host_of(url: str) -> str:
        if not url:
            return ""
        raw = url.strip()
        parsed = urlparse(raw if "://" in raw else "//" + raw)
        host = (parsed.hostname or "").lower()
        if host.startswith("www."):
            host = host[4:]
        return host

    def _is_host(host: str, domain: str) -> bool:
        if not host:
            return False
        return host == domain or host.endswith("." + domain)

    def classify_website(url: str) -> str:
        if not url or url.strip().lower() in ["-", "n/a", "none", "null"]:
            return "NO_WEBSITE"
        url_lower = url.lower()
        host = _host_of(url)

        social_domains = [
            "instagram.com", "facebook.com", "fb.com", "linkedin.com", "t.me",
            "telegram.me", "threads.net", "youtube.com", "pinterest.com",
            "twitter.com", "x.com", "tiktok.com", "wa.me", "whatsapp.com",
            "linktr.ee", "sites.google.com", "drive.google.com", "business.site"
        ]
        if any(_is_host(host, d) for d in social_domains):
            return "SOCIAL_PROFILE"
        google_platforms = ["google.com/maps", "business.google.com", "g.page"]
        if any(_is_host(host, p.split("/")[0]) and p in url_lower for p in google_platforms):
            return "GOOGLE_BUSINESS"
        marketplace_platforms = ["yelp.com", "tripadvisor.com", "amazon.com", "ebay.com"]
        if any(_is_host(host, d) for d in marketplace_platforms):
            return "MARKETPLACE_PROFILE"
        return "PROFESSIONAL_WEBSITE"

    web_type = classify_website(website)
    
    if web_type != "PROFESSIONAL_WEBSITE":
        # Skip analysis and update the lead
        update_data = {
            "website_type": web_type,
            "social_only": True,
            "analysis_skipped_reason": "NO_PROFESSIONAL_WEBSITE",
            "ai_analysis.analysis_status": "SKIPPED",
            "ai_analysis.analysis_completed_at": datetime.utcnow()
        }
        await db["leads"].update_one(query, {"$set": update_data})
        
        # We return empty IDs. Frontend should handle empty queue_id.
        return StartAnalysisResponse(
            lead_id=str(id),
            website_url=website,
            queue_id="",
            worker_id="skipped"
        )

    if not website.startswith('http'):
        website = "http://" + website

    from datetime import timezone
    # Check for genuinely active jobs; if stale (>300s), fail it so a fresh analysis runs
    now = datetime.now(timezone.utc)
    active_job = await db["ai_analysis_jobs"].find_one({
        "lead_id": str(id),
        "status": {"$in": ["WAITING", "STARTED", "COLLECTING", "ANALYZING", "SCORING", "SAVING", "QUEUED"]}
    }, sort=[("_id", -1)])
    if active_job:
        job_updated = active_job.get("updated_at") or active_job.get("started_at") or active_job.get("created_at")
        if job_updated and job_updated.tzinfo is None:
            job_updated = job_updated.replace(tzinfo=timezone.utc)
        age = (now - job_updated).total_seconds() if job_updated else 999
        if age <= 300:
            return StartAnalysisResponse(
                lead_id=str(id),
                website_url=website,
                queue_id=str(active_job.get("queue_id", "")),
                worker_id="Worker-Active"
            )
        else:
            # Stale job auto-failed
            await db["ai_analysis_jobs"].update_one(
                {"_id": active_job["_id"]},
                {"$set": {"status": "FAILED", "error_message": "Stale job superseded by fresh analysis request.", "updated_at": now}}
            )

    from app.services.ai_analysis_queue_service import AIAnalysisQueueService
    
    lead_copy = dict(lead)
    lead_copy["website"] = website
    
    queue = await AIAnalysisQueueService.create_queue(
        user_id=str(current_user.id),
        leads=[lead_copy]
    )
    
    # Immediately update the lead document with genuine processing state
    now = datetime.now(timezone.utc)
    await db["leads"].update_one(
        {"_id": lead["_id"]},
        {"$set": {
            "ai_analysis.analysis_status": "Queued",
            "ai_analysis.progress": 10,
            "ai_analysis.stage": "Queued for AI analysis",
            "ai_analysis.queue_id": str(queue.id),
            "ai_analysis.analysis_started_at": now,
            "updated_at": now,
        }}
    )
    
    return StartAnalysisResponse(
        lead_id=str(id),
        website_url=website,
        queue_id=str(queue.id),
        worker_id="Worker-Alpha"
    )

