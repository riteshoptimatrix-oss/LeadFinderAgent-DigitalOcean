"""Excel Report Generator"""
import os
from typing import Dict, Any
import logging

from app.models.report import ReportType

logger = logging.getLogger(__name__)

try:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    PANDAS_AVAILABLE = True
except Exception as exc:
    pd = None
    load_workbook = None
    Font = PatternFill = Alignment = None
    PANDAS_AVAILABLE = False
    logger.warning("[EXCEL] Excel report dependencies unavailable: %s", exc)

class ExcelReportGenerator:
    def __init__(self):
        self.reports_dir = "reports/excel"
        os.makedirs(self.reports_dir, exist_ok=True)
    
    async def generate(self, report_id: str, report_data: Dict[str, Any],
                      report_type: ReportType) -> tuple[str, int]:
        if not PANDAS_AVAILABLE or pd is None or load_workbook is None:
            raise RuntimeError("Excel report generation is unavailable because pandas/openpyxl dependencies are not installed or are broken")

        filename = f"{report_id}_{report_type.value.replace(' ', '_')}.xlsx"
        file_path = os.path.join(self.reports_dir, filename)
        
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            # Lead Info Sheet
            lead = report_data.get('lead', {})
            lead_df = pd.DataFrame([{
                'Business Name': lead.get('name', 'N/A'),
                'Website': lead.get('website', 'N/A'),
                'Email': lead.get('email', 'N/A'),
                'Phone': lead.get('phone', 'N/A'),
                'City': lead.get('city', 'N/A'),
                'State': lead.get('state', 'N/A'),
                'Country': lead.get('country', 'N/A')
            }])
            lead_df.to_excel(writer, sheet_name='Lead Info', index=False)
            
            # Scores Sheet
            ai = report_data.get('ai_analysis', {})
            scores_df = pd.DataFrame([{
                'Overall Score': ai.get('overall_ai_score', 0),
                'SEO Score': ai.get('seo_score', 0),
                'Performance Score': ai.get('performance_score', 0),
                'Security Score': ai.get('security_score', 0),
                'UI Score': ai.get('ui_score', 0),
                'UX Score': ai.get('ux_score', 0),
                'Business Opportunity': ai.get('business_opportunity_score', 0)
            }])
            scores_df.to_excel(writer, sheet_name='Scores', index=False)
            
            # SEO Analysis
            seo = report_data.get('seo_analysis', {})
            if seo and seo.get('issues'):
                seo_issues = []
                for issue in seo['issues']:
                    seo_issues.append({
                        'Severity': issue.get('severity', 'N/A'),
                        'Category': issue.get('category', 'N/A'),
                        'Issue Type': issue.get('issue_type', 'N/A'),
                        'Description': issue.get('description', 'N/A')
                    })
                seo_df = pd.DataFrame(seo_issues)
                seo_df.to_excel(writer, sheet_name='SEO Issues', index=False)
            
            # Performance Analysis
            perf = report_data.get('performance_analysis', {})
            if perf and perf.get('issues'):
                perf_issues = []
                for issue in perf['issues']:
                    perf_issues.append({
                        'Issue Type': issue.get('issue_type', 'N/A'),
                        'Priority': issue.get('priority', 'N/A'),
                        'Problem': issue.get('problem', 'N/A'),
                        'Impact': issue.get('impact', 'N/A')
                    })
                perf_df = pd.DataFrame(perf_issues)
                perf_df.to_excel(writer, sheet_name='Performance Issues', index=False)

            # Website Performance Audit
            perf_audit = report_data.get('performance_audit', {})
            if perf_audit:
                audit_rows = [{
                    'Metric': 'Overall Score',
                    'Value': perf_audit.get('overallScore', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Grade',
                    'Value': perf_audit.get('grade', 'N/A'),
                    'Category': 'Overall',
                }]
                for section_name, section in (
                    ('coreWebVitals', 'Core Web Vitals'),
                    ('websiteSpeed', 'Website Speed'),
                ):
                    data = perf_audit.get(section_name, {}) or {}
                    if not isinstance(data, dict):
                        continue
                    for key, value in data.items():
                        if isinstance(value, dict):
                            label = value.get('label')
                            if not label and (value.get('value') is not None or value.get('unit')):
                                label = f"{value.get('value', '')} {value.get('unit', '')}".strip()
                            audit_rows.append({
                                'Metric': f"{section_name} {key}",
                                'Value': label or 'N/A',
                                'Category': section_name,
                            })
                for key in ('images', 'javascript', 'css', 'fonts', 'compression', 'cache', 'cdn', 'resources', 'mobile', 'desktop', 'hosting', 'thirdParty'):
                    data = perf_audit.get(key, {}) or {}
                    if isinstance(data, dict):
                        audit_rows.append({'Metric': key, 'Value': data.get('score', 'N/A'), 'Category': key})

                audit_df = pd.DataFrame(audit_rows)
                audit_df.to_excel(writer, sheet_name='Performance Audit', index=False)

                recommendations = perf_audit.get('recommendations', []) or []
                if recommendations:
                    rec_rows = []
                    for rec in recommendations:
                        rec_rows.append({
                            'Title': rec.get('title', 'N/A'),
                            'Description': rec.get('description', 'N/A'),
                            'Priority': rec.get('priority', 'N/A'),
                            'Category': rec.get('category', 'N/A'),
                        })
                    rec_df = pd.DataFrame(rec_rows)
                    rec_df.to_excel(writer, sheet_name='Audit Recommendations', index=False)

            # Website Security Audit
            sec_audit = report_data.get('security_audit', {})
            if sec_audit:
                sec_rows = [{
                    'Metric': 'Overall Score',
                    'Value': sec_audit.get('overallScore', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Grade',
                    'Value': sec_audit.get('grade', 'N/A'),
                    'Category': 'Overall',
                }]
                for key in ('ssl', 'headers', 'cookies', 'server', 'forms', 'content', 'authentication', 'domain'):
                    data = sec_audit.get(key, {}) or {}
                    if isinstance(data, dict):
                        sec_rows.append({'Metric': key, 'Value': data.get('score', 'N/A'), 'Category': key})
                        for field, value in data.items():
                            if field == 'score':
                                continue
                            if isinstance(value, list):
                                continue
                            sec_rows.append({'Metric': f"{key} {field}", 'Value': value if value is not None else 'N/A', 'Category': key})

                sec_df = pd.DataFrame(sec_rows)
                sec_df.to_excel(writer, sheet_name='Security Audit', index=False)

                sec_recommendations = sec_audit.get('recommendations', []) or []
                if sec_recommendations:
                    sec_rec_rows = []
                    for rec in sec_recommendations:
                        sec_rec_rows.append({
                            'Title': rec.get('title', 'N/A'),
                            'Description': rec.get('description', 'N/A'),
                            'Priority': rec.get('priority', 'N/A'),
                            'Category': rec.get('category', 'N/A'),
                        })
                    sec_rec_df = pd.DataFrame(sec_rec_rows)
                    sec_rec_df.to_excel(writer, sheet_name='Security Recommendations', index=False)

            # SEO Intelligence Audit
            seo_audit = report_data.get('seo_audit', {})
            if seo_audit:
                seo_rows = [{
                    'Metric': 'Overall Score',
                    'Value': seo_audit.get('overallScore', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Grade',
                    'Value': seo_audit.get('grade', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Status',
                    'Value': seo_audit.get('overallStatus', 'N/A'),
                    'Category': 'Overall',
                }]
                module_map = {
                    'meta': 'Meta Tags',
                    'headings': 'Headings',
                    'images': 'Images',
                    'urls': 'URL Structure',
                    'indexability': 'Indexability',
                    'openGraph': 'Open Graph',
                    'twitterCards': 'Twitter Cards',
                    'schema': 'Schema Markup',
                    'internalLinks': 'Internal Links',
                    'externalLinks': 'External Links',
                    'content': 'Content Quality',
                    'mobile': 'Mobile SEO',
                    'pages': 'Pages',
                    'localSEO': 'Local SEO',
                    'socialSEO': 'Social SEO',
                    'favicon': 'Favicon',
                }
                for key, label in module_map.items():
                    data = seo_audit.get(key, {}) or {}
                    if isinstance(data, dict):
                        seo_rows.append({'Metric': label, 'Value': data.get('score', 'N/A'), 'Category': key})

                seo_df = pd.DataFrame(seo_rows)
                seo_df.to_excel(writer, sheet_name='SEO Audit', index=False)

                seo_issues = seo_audit.get('issues', []) or []
                if seo_issues:
                    seo_issue_rows = []
                    for issue in seo_issues:
                        seo_issue_rows.append({
                            'Title': issue.get('title', 'N/A'),
                            'Description': issue.get('description', 'N/A'),
                            'Severity': issue.get('severity', 'N/A'),
                            'Category': issue.get('category', 'N/A'),
                            'URL': issue.get('affectedUrl', 'N/A'),
                        })
                    seo_issue_df = pd.DataFrame(seo_issue_rows)
                    seo_issue_df.to_excel(writer, sheet_name='SEO Issues', index=False)

                seo_recommendations = seo_audit.get('recommendations', []) or []
                if seo_recommendations:
                    seo_rec_rows = []
                    for rec in seo_recommendations:
                        seo_rec_rows.append({
                            'Title': rec.get('title', 'N/A'),
                            'Description': rec.get('description', 'N/A'),
                            'Priority': rec.get('priority', 'N/A'),
                            'Category': rec.get('category', 'N/A'),
                        })
                    seo_rec_df = pd.DataFrame(seo_rec_rows)
                    seo_rec_df.to_excel(writer, sheet_name='SEO Recommendations', index=False)

            # Website UI/UX Intelligence Audit
            uiux_audit = report_data.get('uiux_audit', {})
            if uiux_audit:
                uiux_rows = [{
                    'Metric': 'Overall Score',
                    'Value': uiux_audit.get('overallScore', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Grade',
                    'Value': uiux_audit.get('grade', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Status',
                    'Value': uiux_audit.get('overallStatus', 'N/A'),
                    'Category': 'Overall',
                }, {
                    'Metric': 'Audit Time (ms)',
                    'Value': uiux_audit.get('auditTime', 'N/A'),
                    'Category': 'Overall',
                }]

                redesign = uiux_audit.get('redesignOpportunity', {}) or {}
                if isinstance(redesign, dict):
                    uiux_rows.append({'Metric': 'Redesign Score', 'Value': redesign.get('score', 'N/A'), 'Category': 'Redesign Opportunity'})
                    uiux_rows.append({'Metric': 'Redesign Level', 'Value': redesign.get('opportunityLevel', 'N/A'), 'Category': 'Redesign Opportunity'})
                    uiux_rows.append({'Metric': 'Redesign Rationale', 'Value': redesign.get('rationale', 'N/A'), 'Category': 'Redesign Opportunity'})
                    uiux_rows.append({'Metric': 'Redesign Estimated Impact', 'Value': redesign.get('estimatedImpact', 'N/A'), 'Category': 'Redesign Opportunity'})
                    uiux_rows.append({'Metric': 'Redesign Payoff', 'Value': redesign.get('payoff', 'N/A'), 'Category': 'Redesign Opportunity'})

                sales_opp = uiux_audit.get('salesOpportunity', {}) or {}
                if isinstance(sales_opp, dict):
                    uiux_rows.append({'Metric': 'Sales Score', 'Value': sales_opp.get('score', 'N/A'), 'Category': 'Sales Opportunity'})
                    uiux_rows.append({'Metric': 'Sales Level', 'Value': sales_opp.get('opportunityLevel', 'N/A'), 'Category': 'Sales Opportunity'})
                    uiux_rows.append({'Metric': 'Pitch Summary', 'Value': sales_opp.get('pitchSummary', 'N/A'), 'Category': 'Sales Opportunity'})

                module_map = {
                    'visualDesign': 'Visual Design',
                    'colors': 'Colors',
                    'typography': 'Typography',
                    'layout': 'Layout',
                    'navigation': 'Navigation',
                    'cta': 'Call-to-Action',
                    'responsive': 'Responsive',
                    'mobile': 'Mobile',
                    'forms': 'Forms',
                    'ux': 'UX',
                    'content': 'Content',
                    'images': 'Images',
                    'icons': 'Icons',
                    'branding': 'Branding',
                    'accessibility': 'Accessibility',
                    'modernDesign': 'Modern Design',
                    'trust': 'Trust',
                }
                for key, label in module_map.items():
                    data = uiux_audit.get(key, {}) or {}
                    if isinstance(data, dict):
                        uiux_rows.append({'Metric': label, 'Value': data.get('score', 'N/A'), 'Category': key})

                uiux_df = pd.DataFrame(uiux_rows)
                uiux_df.to_excel(writer, sheet_name='UI/UX Intelligence', index=False)

                uiux_issues = uiux_audit.get('issues', []) or []
                if uiux_issues:
                    uiux_issue_rows = []
                    for issue in uiux_issues:
                        uiux_issue_rows.append({
                            'Title': issue.get('title', 'N/A'),
                            'Description': issue.get('description', 'N/A'),
                            'Severity': issue.get('severity', 'N/A'),
                            'Category': issue.get('category', 'N/A'),
                            'URL': issue.get('affectedUrl', 'N/A'),
                        })
                    uiux_issue_df = pd.DataFrame(uiux_issue_rows)
                    uiux_issue_df.to_excel(writer, sheet_name='UI/UX Issues', index=False)

                uiux_recommendations = uiux_audit.get('recommendations', []) or []
                if uiux_recommendations:
                    uiux_rec_rows = []
                    for rec in uiux_recommendations:
                        uiux_rec_rows.append({
                            'Title': rec.get('title', 'N/A'),
                            'Description': rec.get('description', 'N/A'),
                            'Priority': rec.get('priority', 'N/A'),
                            'Category': rec.get('category', 'N/A'),
                        })
                    uiux_rec_df = pd.DataFrame(uiux_rec_rows)
                    uiux_rec_df.to_excel(writer, sheet_name='UI/UX Recommendations', index=False)

            # Technologies
            tech = report_data.get('technology_detection', {})
            if tech and tech.get('technologies'):
                tech_list = []
                for t in tech['technologies']:
                    tech_list.append({
                        'Technology': t.get('name', 'N/A'),
                        'Category': t.get('category', 'N/A'),
                        'Version': t.get('version', 'N/A'),
                        'Confidence': f"{t.get('confidence_score', 0)}%"
                    })
                tech_df = pd.DataFrame(tech_list)
                tech_df.to_excel(writer, sheet_name='Technologies', index=False)
            
            # Contacts
            contacts = report_data.get('contact_discovery', {})
            if contacts and contacts.get('contacts'):
                contact_list = []
                for c in contacts['contacts']:
                    contact_list.append({
                        'Type': c.get('type', 'N/A'),
                        'Value': c.get('value', 'N/A'),
                        'Classification': c.get('classification', 'N/A'),
                        'Is Primary': 'Yes' if c.get('is_primary') else 'No'
                    })
                contact_df = pd.DataFrame(contact_list)
                contact_df.to_excel(writer, sheet_name='Contacts', index=False)
        
        # Apply styling
        self._apply_styling(file_path)
        
        file_size = os.path.getsize(file_path)
        return file_path, file_size
    
    def _apply_styling(self, file_path: str):
        wb = load_workbook(file_path)
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(file_path)
